"""Generate safe sample files: no uploaded documents or patient records."""

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
lines = [
    "INBAOCAO — MẪU WORD KEYWORD KHỔ A3",
    "Mẫu minh họa kỹ thuật, không thay thế biểu mẫu y tế theo quy định.",
    "",
    "Họ và tên: $hovaten                         STT: $so",
    "Giới tính: $gt                         Sinh ngày: $ngaysinh",
    "Mã nhân viên: $manv",
    "Số CMND/CCCD /Hộ chiếu/định danh CD: $cccd",
    "Cấp ngày: $ngaycap                         Tại: $noicap",
    "Dân tộc: $dantoc                         Đối tượng: $doituong",
    "Nguồn chi trả: $nguonchitra                         Nhóm máu: $nhommau",
    "Nơi ở hiện tại: $noiohientai",
    "Nghề nghiệp: $nghenghiep",
    "Nơi công tác, học tập: $noicongtac",
    "Lý do khám: $lydokham",
    "Số điện thoại: $sdt                         Bộ phận: $bophan",
    "Chức vụ: $chucvu",
    "",
    "Người đề nghị khám sức khỏe",
    "(Ký và ghi rõ họ tên)",
    "$hovaten",
]
paragraphs = "".join(
    '<w:p><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="26"/></w:rPr><w:t xml:space="preserve">'
    + escape(line)
    + "</w:t></w:r></w:p>"
    for line in lines
)
document = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="{W}"><w:body>{paragraphs}<w:sectPr><w:pgSz w:w="23814" w:h="16840" w:orient="landscape"/><w:pgMar w:top="850" w:right="850" w:bottom="850" w:left="850"/></w:sectPr></w:body></w:document>"""
with ZipFile(ROOT / "MAU_WORD_KEYWORD.docx", "w", ZIP_DEFLATED) as archive:
    archive.writestr(
        "[Content_Types].xml",
        """<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>""",
    )
    archive.writestr(
        "_rels/.rels",
        """<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>""",
    )
    archive.writestr("word/document.xml", document)
book = Workbook()
sheet = book.active
sheet.title = "Danh sách"
sheet.append(["DANH SÁCH MINH HỌA — DỮ LIỆU GIẢ ĐỂ KIỂM THỬ"])
sheet.append([])
headers = [
    "STT",
    "HỌ VÀ TÊN (*)",
    "Giới tính",
    "Mã NV",
    "NGÀY SINH",
    "CCCD",
    "Ngày cấp CCCD",
    "Nơi cấp CCCD",
    "Dân tộc",
    "Nơi ở hiện tại",
    "Nghề nghiệp",
    "Công ty",
    "Đối tượng",
    "Nguồn chi trả",
    "Nhóm máu",
    "Lý do khám",
    "SĐT",
    "Bộ phận",
    "Chức vụ",
]
sheet.append(headers)
sheet.append(
    [
        1,
        "NGƯỜI THỬ NGHIỆM A",
        "Nữ",
        "NV001",
        datetime(2000, 2, 3),
        "001234567890",
        datetime(2025, 1, 1),
        "Bộ Công an",
        "Kinh",
        "Địa chỉ minh họa",
        "Kế toán",
        "Đơn vị minh họa",
        "Người lao động",
        "Đơn vị",
        "O",
        "Khám định kỳ",
        "0900000000",
        "Văn phòng",
        "Nhân viên",
    ]
)
sheet.append(
    [
        2,
        "NGƯỜI THỬ NGHIỆM B",
        "Nam",
        "NV002",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "Đơn vị minh họa",
    ]
)
for cell in sheet[3]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="147C64")
    cell.alignment = Alignment(wrap_text=True)
sheet.row_dimensions[3].height = 40
for col in sheet.columns:
    sheet.column_dimensions[col[0].column_letter].width = 23
for row in sheet.iter_rows(min_row=4):
    for cell in row:
        if isinstance(cell.value, datetime):
            cell.number_format = "dd/mm/yyyy"
sheet.freeze_panes = "C4"
sheet.auto_filter.ref = f"A3:S{sheet.max_row}"
book.save(ROOT / "DANH_SACH_MAU.xlsx")
print("Generated Word keyword template and synthetic Excel sample.")

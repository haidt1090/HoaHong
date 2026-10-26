from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from datetime import datetime
import json
import threading
import unittest
from openpyxl import Workbook
from inbaocao.core import (
    ExcelData,
    auto_mapping,
    merge_docx,
    scan_tokens,
    row_values,
    W,
)
from inbaocao.storage import Library
from inbaocao.printing import print_batch


def make_docx(path):
    main = f"""<w:document xmlns:w="{W}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14"><w:body>
<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Họ và tên: $ho</w:t></w:r><w:r><w:t>vaten; Giới tính: $gt; $missing</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>${{Tên bổ sung}} / $cccd / $ngaysinh</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
<w:sectPr><w:pgSz w:w="23814" w:h="16840" w:orient="landscape"/></w:sectPr></w:body></w:document>"""
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", main)
        archive.writestr(
            "word/header1.xml",
            f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>$hovaten $manv</w:t></w:r></w:p></w:hdr>',
        )
        archive.writestr(
            "word/footer1.xml",
            f'<w:ftr xmlns:w="{W}"><w:p><w:r><w:t>$hovaten</w:t></w:r></w:p></w:ftr>',
        )
        archive.writestr("word/styles.xml", "<styles>unchanged</styles>")
        archive.writestr("word/media/image1.png", b"opaque original image bytes")


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.work = TemporaryDirectory()
        self.root = Path(self.work.name)
        self.template = self.root / "template.docx"
        make_docx(self.template)

    def tearDown(self):
        self.work.cleanup()

    def test_split_run_keywords_headers_tables_footer_and_preserved_layout(self):
        tokens = scan_tokens(self.template)
        self.assertEqual(tokens.count("$hovaten"), 1)
        self.assertIn("${Tên bổ sung}", tokens)
        output = self.root / "filled.docx"
        merge_docx(
            self.template,
            output,
            {
                "$hovaten": "Đặng Thị Ánh & <An>",
                "$gt": "Nữ",
                "${Tên bổ sung}": "Điều dưỡng",
                "$cccd": "001234567890",
            },
        )
        with ZipFile(output) as archive:
            xml = archive.read("word/document.xml")
            self.assertIn(b"xmlns:w14=", xml)
            self.assertIn(b'mc:Ignorable="w14"', xml)
            root = ET.fromstring(xml)
            texts = "".join(n.text or "" for n in root.iter(f"{{{W}}}t"))
            self.assertIn("Đặng Thị Ánh & <An>", texts)
            self.assertNotIn("$", texts)
            self.assertEqual(root.find(f".//{{{W}}}pgSz").get(f"{{{W}}}w"), "23814")
            self.assertIsNotNone(root.find(f".//{{{W}}}b"))
            self.assertEqual(
                archive.read("word/styles.xml"), b"<styles>unchanged</styles>"
            )
            self.assertEqual(
                archive.read("word/media/image1.png"), b"opaque original image bytes"
            )
            for name in ["word/header1.xml", "word/footer1.xml"]:
                self.assertIn(
                    "Đặng Thị Ánh",
                    "".join(ET.fromstring(archive.read(name)).itertext()),
                )
        self.assertIn("$hovaten", scan_tokens(self.template))
        self.assertEqual(scan_tokens(output), [])

    def test_excel_header_dates_leading_zeros_optional_and_custom_keywords(self):
        filename = self.root / "list.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Danh sách"
        sheet.append(["DANH SÁCH"])
        sheet.append([])
        sheet.append(
            [
                "STT",
                "HỌ VÀ TÊN (*)",
                "gioi tinh",
                "MÃ NV(*)",
                "NGÀY SINH",
                "Số CCCD",
                "Tên bổ sung",
            ]
        )
        sheet.append(
            [1, "Nguyễn Ánh", "Nữ", "NV001", datetime(2000, 2, 3), 12345678, "Kỹ thuật"]
        )
        sheet.cell(4, 6).number_format = "000000000000"
        book.save(filename)
        data = ExcelData(filename)
        try:
            self.assertEqual(data.guess_header("Danh sách"), 3)
            headers, rows = data.read("Danh sách", 3)
        finally:
            data.close()
        tokens = scan_tokens(self.template)
        mapping = auto_mapping(tokens, headers)
        self.assertEqual(mapping["$hovaten"], "HỌ VÀ TÊN (*)")
        self.assertEqual(mapping["$gt"], "gioi tinh")
        values = row_values(tokens, rows[0], mapping, {"$missing": "Công ty chung"})
        self.assertEqual(values["$ngaysinh"], "03/02/2000")
        self.assertEqual(values["$cccd"], "000012345678")
        self.assertEqual(values["${Tên bổ sung}"], "Kỹ thuật")
        self.assertEqual(values["$missing"], "Công ty chung")
        self.assertEqual(row_values(["$x"], {}, {}), {"$x": ""})

    def test_library_persistence_backup_restore_and_no_patient_list(self):
        library = Library(self.root / "one")
        item = library.add(self.template, "Mẫu A3")
        item["mapping"] = {"$hovaten": "HỌ VÀ TÊN"}
        library.save()
        self.assertEqual(
            Library(self.root / "one").templates[0]["mapping"], item["mapping"]
        )
        backup = self.root / "backup.zip"
        library.backup(backup)
        other = Library(self.root / "two")
        self.assertEqual(other.restore(backup), 1)
        self.assertEqual(other.templates[0]["name"], "Mẫu A3")
        self.assertNotEqual(other.templates[0]["id"], item["id"])
        self.assertEqual(
            scan_tokens(other.file(other.templates[0])), scan_tokens(self.template)
        )
        with ZipFile(backup) as archive:
            self.assertEqual(len(archive.namelist()), 2)

    def test_print_batch_one_job_per_person_no_review_and_cleanup(self):
        class Document:
            def __init__(self):
                self.calls = []

            def PrintOut(self, **kwargs):
                self.calls.append(("print", kwargs))

            def Close(self, **kwargs):
                self.calls.append(("close", kwargs))

        class Documents:
            def __init__(self):
                self.docs = []

            def Open(self, path, **kwargs):
                self_outer.assertEqual(scan_tokens(path), [])
                self_outer.assertTrue(kwargs["ReadOnly"])
                doc = Document()
                self.docs.append(doc)
                return doc

        class Word:
            def __init__(self):
                self.Documents = Documents()
                self.quit = False
                self.WordBasic = self
                # Reproduce WordBasic's missing method metadata on Windows.
                self.FilePrintSetup = None
                self.flagged = []

            def _FlagAsMethod(self, name):
                self.flagged.append(name)
                self.FilePrintSetup = self.setup_printer

            def setup_printer(self, **kwargs):
                self.printer = kwargs

            def Quit(self, **kwargs):
                self.quit = True

        self_outer = self
        word = Word()
        progress = []
        total = print_batch(
            self.template,
            scan_tokens(self.template),
            [{"name": "An"}, {"name": "Bình"}],
            {"$hovaten": "name"},
            {},
            "Máy in A3",
            threading.Event(),
            lambda a, b: progress.append((a, b)),
            dispatch=lambda _: word,
        )
        self.assertEqual(total, 2)
        self.assertEqual(progress, [(1, 2), (2, 2)])
        self.assertTrue(word.quit)
        self.assertFalse(word.Visible)
        self.assertEqual(word.flagged, ["FilePrintSetup"])
        self.assertEqual(word.printer["DoNotSetAsSysDefault"], 1)
        for doc in word.Documents.docs:
            self.assertEqual(
                doc.calls[0], ("print", {"Background": False, "Copies": 1})
            )
            self.assertEqual(doc.calls[1][0], "close")

    def test_print_cancel_between_jobs_and_failure_quits_word(self):
        event = threading.Event()

        class Doc:
            def PrintOut(self, **kwargs):
                event.set()

            def Close(self, **kwargs):
                pass

        class Word:
            Visible = True
            quit = False

            @property
            def Documents(self):
                return self

            def Open(self, *args, **kwargs):
                return Doc()

            def Quit(self, **kwargs):
                self.quit = True

        word = Word()
        total = print_batch(
            self.template,
            scan_tokens(self.template),
            [{}, {}],
            {},
            {},
            "",
            event,
            lambda *_: None,
            dispatch=lambda _: word,
        )
        self.assertEqual(total, 1)
        self.assertTrue(word.quit)

        class Broken(Word):
            def Open(self, *args, **kwargs):
                raise RuntimeError("Printer error")

        broken = Broken()
        with self.assertRaisesRegex(RuntimeError, "Printer error"):
            print_batch(
                self.template,
                scan_tokens(self.template),
                [{}],
                {},
                {},
                "",
                threading.Event(),
                lambda *_: None,
                dispatch=lambda _: broken,
            )
        self.assertTrue(broken.quit)

    def test_print_error_identifies_stage_and_preserves_original_error(self):
        class Word:
            WordBasic = None

            def Quit(self, **kwargs):
                raise RuntimeError("cleanup failed")

        with self.assertRaisesRegex(RuntimeError, "chọn máy in Máy A3") as caught:
            print_batch(
                self.template, [], [{}], {}, {}, "Máy A3",
                threading.Event(), lambda *_: None, dispatch=lambda _: Word(),
            )
        self.assertNotIn("cleanup failed", str(caught.exception))

    def test_restore_does_not_extract_paths_outside_store(self):
        backup = self.root / "bad.zip"
        with ZipFile(backup, "w") as archive:
            archive.writestr(
                "library.json",
                json.dumps(
                    {
                        "format": "inbaocao-word-v2",
                        "templates": [{"id": "../../escape", "name": "bad"}],
                    }
                ),
            )
            archive.writestr("../../escape.docx", "bad")
        with self.assertRaises(ValueError):
            Library(self.root / "library").restore(backup)
        self.assertFalse((self.root / "escape.docx").exists())


if __name__ == "__main__":
    unittest.main()

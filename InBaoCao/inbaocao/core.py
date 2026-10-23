"""Excel reading and DOCX replacement; no Word installation needed for merging."""

from __future__ import annotations
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
TEXT = f"{{{W}}}t"
PARA = f"{{{W}}}p"
TOKEN = re.compile(r"\$(?:\{([^{}\r\n]+)\}|([^\W\d]\w*))", re.UNICODE)
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
# Recommended keywords and equivalent Excel headers. Any additional keyword can
# be mapped by the user; these aliases are only defaults, not required columns.
FIELDS = [
    ("so", "Số / STT", ["STT", "Số"]),
    ("hovaten", "Họ và tên", ["Họ tên", "Họ và tên"]),
    ("gt", "Giới tính", ["Giới tính"]),
    ("manv", "Mã nhân viên", ["Mã NV", "Mã nhân viên"]),
    ("ngaysinh", "Sinh ngày", ["Ngày sinh", "Sinh ngày"]),
    (
        "cccd",
        "Số CMND/CCCD /Hộ chiếu/định danh CD",
        ["CCCD", "Số CCCD", "CMND", "Hộ chiếu", "Định danh CD"],
    ),
    ("ngaycap", "Cấp ngày", ["Ngày cấp CCCD", "Ngày cấp", "Cấp ngày"]),
    ("noicap", "Tại", ["Nơi cấp CCCD", "Nơi cấp"]),
    ("dantoc", "Dân tộc", ["Dân tộc"]),
    ("noiohientai", "Nơi ở hiện tại", ["Nơi ở hiện tại"]),
    ("nghenghiep", "Nghề nghiệp", ["Nghề nghiệp"]),
    (
        "noicongtac",
        "Nơi công tác, học tập",
        ["Nơi công tác, học tập", "Nơi công tác", "Công ty", "Đơn vị", "Nơi làm việc"],
    ),
    ("doituong", "Đối tượng", ["Đối tượng"]),
    ("nguonchitra", "Nguồn chi trả", ["Nguồn chi trả"]),
    ("nhommau", "Nhóm máu", ["Nhóm máu"]),
    ("lydokham", "Lý do khám", ["Lý do khám", "Lí do khám"]),
    ("sdt", "Số điện thoại", ["SĐT", "Số điện thoại"]),
    ("bophan", "Bộ phận", ["Bộ phận"]),
    ("chucvu", "Chức vụ", ["Chức vụ", "Chưc vụ"]),
]
EXTRA_ALIASES = {
    "hoten": "hovaten",
    "gioitinh": "gt",
    "stt": "so",
    "sinhngay": "ngaysinh",
    "ngaycapcccd": "ngaycap",
    "noicapcccd": "noicap",
    "diachihientai": "noiohientai",
    "manhanvien": "manv",
    "congty": "noicongtac",
    "noicongtachoctap": "noicongtac",
}


def normalize(value) -> str:
    value = str(value or "").replace("Đ", "D").replace("đ", "d")
    return "".join(
        c
        for c in unicodedata.normalize("NFD", value).lower()
        if c.isascii() and c.isalnum()
    )


def token_key(token: str) -> str:
    return normalize(token[2:-1] if token.startswith("${") else token.lstrip("$"))


def aliases_for(token: str) -> list[str]:
    key = token_key(token)
    key = EXTRA_ALIASES.get(key, key)
    for name, label, aliases in FIELDS:
        if key in {normalize(name), normalize(label), *(normalize(a) for a in aliases)}:
            return [name, label, *aliases]
    return [key]


def auto_mapping(tokens: list[str], headers: list[str]) -> dict[str, str]:
    return {
        t: next(
            (
                h
                for h in headers
                if normalize(h) in {normalize(a) for a in aliases_for(t)}
            ),
            "",
        )
        for t in tokens
    }


def cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, (int, float)) and value == int(value):
        text = str(int(value))
        fmt = cell.number_format.split(";")[0]
        if re.fullmatch(r"0+", fmt):
            return text.zfill(len(fmt))
        return text
    return str(value)


class ExcelData:
    def __init__(self, filename):
        self.book = load_workbook(filename, data_only=True, read_only=True)
        self.sheets = self.book.sheetnames

    def close(self):
        self.book.close()

    def guess_header(self, sheet_name):
        best, score = 1, -1
        known = {
            normalize(a)
            for key, label, aliases in FIELDS
            for a in [key, label, *aliases]
        }
        for index, row in enumerate(self.book[sheet_name].iter_rows(max_row=40), 1):
            names = [normalize(c.value) for c in row if c.value is not None]
            points = sum(n in known for n in names) + 2 * ("stt" in names)
            if points > score:
                best, score = index, points
        return best

    def read(self, sheet_name, header_row):
        if header_row < 1:
            raise ValueError("Dòng tiêu đề phải từ 1 trở lên.")
        rows = iter(self.book[sheet_name].iter_rows(min_row=header_row))
        first = next(rows, ())
        headers, positions = [], []
        for col, cell in enumerate(first):
            h = cell_text(cell).strip()
            if not h:
                continue
            while h in headers:
                h += " (trùng)"
            headers.append(h)
            positions.append(col)
        if not headers:
            raise ValueError("Dòng được chọn không có tiêu đề.")
        result = []
        for row in rows:
            record = {
                h: cell_text(row[col]) if col < len(row) else ""
                for h, col in zip(headers, positions)
            }
            if any(v.strip() for v in record.values()):
                result.append(record)
        return headers, result


def paragraph_groups(root):
    """Nearest paragraph grouping includes tables, headers and nested text boxes once."""
    parents = {child: parent for parent in root.iter() for child in parent}
    groups = {}
    for node in root.iter(TEXT):
        ancestor = parents.get(node)
        while ancestor is not None and ancestor.tag != PARA:
            ancestor = parents.get(ancestor)
        if ancestor is not None:
            groups.setdefault(ancestor, []).append(node)
    return groups.values()


def _xml_roots(archive):
    for name in archive.namelist():
        if name.startswith("word/") and name.endswith(".xml"):
            data = archive.read(name)
            try:
                root = ET.fromstring(data)
            except ET.ParseError:
                continue
            if any(True for _ in root.iter(TEXT)):
                yield name, data, root


def scan_tokens(filename: str | Path) -> list[str]:
    tokens = []
    with ZipFile(filename) as archive:
        for _, _, root in _xml_roots(archive):
            for group in paragraph_groups(root):
                for match in TOKEN.finditer("".join(n.text or "" for n in group)):
                    if match.group() not in tokens:
                        tokens.append(match.group())
    return tokens


def _serialize(root, original):
    from io import BytesIO
    from xml.sax.saxutils import quoteattr

    namespaces = []
    for event, value in ET.iterparse(BytesIO(original), events=["start-ns", "start"]):
        if event == "start":
            break
        namespaces.append(value)
        try:
            ET.register_namespace(*value)
        except ValueError:
            pass
    output = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    start = output.find(b"<", output.find(b"?>") + 2)
    end = output.find(b">", start)
    opening = output[start:end]
    # ElementTree omits unused declarations, but Word mc:Ignorable may still
    # reference them. Preserve original root declarations for valid DOCX parts.
    additions = []
    for prefix, uri in namespaces:
        attribute = ("xmlns:" + prefix) if prefix else "xmlns"
        if (attribute + "=").encode() not in opening:
            additions.append((" " + attribute + "=" + quoteattr(uri)).encode("utf-8"))
    return output[:end] + b"".join(additions) + output[end:]


def merge_docx(source, destination, values: dict[str, str]):
    """Replace full keywords even if Word split them across differently styled runs.

    Only affected text nodes change. All untouched XML parts remain byte-for-byte;
    section/page sizes, images, tables, header/footer and styles are retained.
    """
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError("Không ghi đè file mẫu gốc.")
    with ZipFile(source) as archive:
        updates = {}
        for name, original, root in _xml_roots(archive):
            changed = False
            for group in paragraph_groups(root):
                text = "".join(n.text or "" for n in group)
                matches = list(TOKEN.finditer(text))
                if not matches:
                    continue
                changed = True
                ranges, offset = [], 0
                for node in group:
                    length = len(node.text or "")
                    ranges.append((node, offset, offset + length))
                    offset += length
                # Right-to-left replacement retains the original offsets.
                for match in reversed(matches):
                    touched = [
                        (n, a, b)
                        for n, a, b in ranges
                        if b > match.start() and a < match.end()
                    ]
                    replacement = (
                        str(values.get(match.group(), "") or "")
                        .replace("\r", " ")
                        .replace("\n", " ")
                    )
                    for index, (node, start, end) in enumerate(touched):
                        current = node.text or ""
                        left, right = max(0, match.start() - start), min(
                            end - start, match.end() - start
                        )
                        node.text = (
                            current[:left]
                            + (replacement if index == 0 else "")
                            + current[right:]
                        )
                        node.set(XML_SPACE, "preserve")
            if changed:
                updates[name] = _serialize(root, original)
        destination.parent.mkdir(parents=True, exist_ok=True)
        # Preserve zip metadata and every unrelated package entry.
        with ZipFile(destination, "w", ZIP_DEFLATED) as output:
            for info in archive.infolist():
                output.writestr(
                    info, updates.get(info.filename, archive.read(info.filename))
                )


def row_values(tokens, row, mapping, constants=None):
    constants = constants or {}
    return {
        t: (
            str(row.get(mapping[t], "") or "")
            if mapping.get(t)
            else str(constants.get(t, "") or "")
        )
        for t in tokens
    }
